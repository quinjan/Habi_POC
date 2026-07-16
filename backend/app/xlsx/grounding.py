from copy import deepcopy
from typing import Literal

from openpyxl.utils import column_index_from_string, get_column_letter

from backend.app.evidence.annotation_policy import is_workflow_noise


_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "line_type": ("line kind", "line type"),
    "name": ("purchase line description", "item", "name"),
    "quantity": ("quantity", "qty"),
    "unit": ("unit",),
    "unit_price": ("unit price",),
    "price": ("combined price", "total price", "price", "unit price"),
    "currency": ("currency",),
    "purchase_date": ("purchase date", "date"),
    "material_name": ("material name",),
    "material_category": ("material category path", "material category"),
    "service_name": ("service name",),
    "service_category": ("service category path", "service category"),
    "provider_state": (
        "reviewer final state",
        "provider state",
        "ai provider state",
    ),
    "provider_name": ("provider name", "supplier name", "supplier", "provider"),
    "provider_category": ("provider category path", "provider category"),
    "remarks_or_terms": ("remarks or terms", "remarks", "terms"),
}

_ANNOTATION_SUFFIX_TYPES: dict[str, str] = {
    "delivery terms": "delivery_terms",
    "payment terms": "payment_terms",
    "validity terms": "validity_terms",
    "validity": "validity_terms",
    "warranty terms": "warranty_terms",
    "warranty": "warranty_terms",
    "availability terms": "availability_terms",
    "availability": "availability_terms",
    "condition or exclusion": "condition_or_exclusion",
    "conditions": "condition_or_exclusion",
    "exclusions": "condition_or_exclusion",
    "remarks or terms": "general_qualifier",
    "remarks": "general_qualifier",
    "terms": "general_qualifier",
    "notes": "general_qualifier",
}


def sanitize_profile_column_mappings(profile: dict, artifact: dict) -> tuple[dict, int]:
    sanitized = deepcopy(profile)
    artifact_columns = {cell["column"] for cell in artifact["cells"]}
    invalid_count = 0
    for region in sanitized.get("regions", []):
        columns = region.get("columns")
        if not isinstance(columns, dict):
            continue
        for field, value in columns.items():
            if value is None:
                continue
            normalized = value.strip().upper() if isinstance(value, str) else ""
            try:
                column_index = column_index_from_string(normalized)
            except ValueError:
                column_index = None
            if column_index is None or column_index not in artifact_columns:
                columns[field] = None
                invalid_count += 1
                continue
            columns[field] = get_column_letter(column_index)
    return sanitized, invalid_count


def ground_profile_to_source(profile: dict, artifact: dict) -> dict:
    grounded = deepcopy(profile)
    artifact_rows = {cell["row"] for cell in artifact["cells"]}
    header_rows = sorted(
        {
            *grounded.get("header_rows", []),
            *(
                row
                for region in grounded.get("regions", [])
                for row in region.get("header_row_numbers", [])
            ),
        }
    )
    if header_rows:
        first_header_row = min(header_rows)
        grounded["title_rows"] = sorted(
            {
                *grounded.get("title_rows", []),
                *(row for row in artifact_rows if row < first_header_row),
            }
        )

    cells_by_row = _cells_by_row_and_column(artifact)
    for region in grounded.get("regions", []):
        columns = dict(region.get("columns", {}))
        header_map = _header_map(
            cells_by_row,
            region.get("header_row_numbers", []) or grounded.get("header_rows", []),
        )
        for field, aliases in _HEADER_ALIASES.items():
            for alias in aliases:
                if alias in header_map:
                    columns[field] = header_map[alias]
                    break
            else:
                columns.setdefault(field, None)
        for header, column in header_map.items():
            annotation = _annotation_header(header)
            if annotation is None:
                continue
            target, annotation_type = annotation
            columns[
                f"annotation__{target}__{annotation_type}__{column}"
            ] = column
        region["columns"] = columns
    return grounded


def build_explicit_row_candidates(
    *,
    rows: list[dict],
    region: dict,
    source_submission_id: int,
    source_file_id: int,
    worksheet_name: str,
) -> dict[int, dict]:
    columns = region.get("columns", {})
    if not columns.get("line_type"):
        return {}

    candidates: dict[int, dict] = {}
    for row in rows:
        candidate = _candidate_from_explicit_row(
            row=row,
            columns=columns,
            source_submission_id=source_submission_id,
            source_file_id=source_file_id,
            worksheet_name=worksheet_name,
            region_id=region["region_id"],
        )
        if candidate is not None:
            candidates[row["row"]] = candidate
    return candidates


def _candidate_from_explicit_row(
    *,
    row: dict,
    columns: dict,
    source_submission_id: int,
    source_file_id: int,
    worksheet_name: str,
    region_id: str,
) -> dict | None:
    cells = {cell["column"]: cell for cell in row["cells"]}
    line_type = _line_type(_cell_text(cells, columns.get("line_type")))
    if line_type is None:
        return None

    linked_concepts: list[dict] = []
    if line_type in {"material", "bundled"}:
        concept = _linked_concept(
            concept_type="material",
            name=_cell_text(cells, columns.get("material_name")),
            category_path=_cell_text(cells, columns.get("material_category")),
        )
        if concept is None:
            return None
        linked_concepts.append(concept)
    if line_type in {"service", "bundled"}:
        concept = _linked_concept(
            concept_type="service",
            name=_cell_text(cells, columns.get("service_name")),
            category_path=_cell_text(cells, columns.get("service_category")),
        )
        if concept is None:
            return None
        linked_concepts.append(concept)

    price = _cell_text(cells, columns.get("price"))
    currency = _cell_text(cells, columns.get("currency"))
    if price is not None and currency is None:
        currency = "PHP"
        currency_state = "defaulted"
    elif currency is not None:
        currency_state = "source_stated"
    else:
        currency_state = "unknown"

    provider_name = _cell_text(cells, columns.get("provider_name"))
    provider_state = _provider_state(
        _cell_text(cells, columns.get("provider_state")), provider_name
    )
    provider_category = (
        _category_suggestion(_cell_text(cells, columns.get("provider_category")))
        if provider_state == "external"
        else None
    )
    if provider_state == "external" and provider_category is None:
        provider_category = {
            "top_level_category": "Providers",
            "subcategory": "General",
        }
    if provider_state == "unknown":
        provider_name = None

    row_number = row["row"]
    annotation_result = _annotation_proposals(
        cells=cells,
        columns=columns,
        worksheet_name=worksheet_name,
        row_number=row_number,
    )
    return {
        "linked_concepts": linked_concepts,
        "quantity": _cell_text(cells, columns.get("quantity")),
        "unit": _cell_text(cells, columns.get("unit")),
        "price": price,
        "currency": currency,
        "currency_state": currency_state,
        "provider_state": provider_state,
        "provider_name": provider_name,
        "provider_category_suggestion": provider_category,
        "purchase_date": _cell_text(cells, columns.get("purchase_date")),
        "remarks_or_terms": _cell_text(cells, columns.get("remarks_or_terms")),
        **annotation_result,
        "confidence": 1.0,
        "evidence": {
            "source_submission_id": source_submission_id,
            "source_file_id": source_file_id,
            "worksheet": worksheet_name,
            "region_id": region_id,
            "primary_body_row": row_number,
            "locators": [{"row": row_number, "role": "body"}],
        },
    }


def _annotation_header(header: str) -> tuple[str, str] | None:
    for target in ("material", "service", "provider"):
        prefix = f"{target} "
        if header.startswith(prefix):
            annotation_type = _ANNOTATION_SUFFIX_TYPES.get(header[len(prefix) :])
            return (target, annotation_type) if annotation_type is not None else None
    annotation_type = _ANNOTATION_SUFFIX_TYPES.get(header)
    return ("purchase_line", annotation_type) if annotation_type is not None else None


def _annotation_proposals(
    *,
    cells: dict[int, dict],
    columns: dict,
    worksheet_name: str,
    row_number: int,
) -> dict:
    mappings: list[tuple[int, str, str, str]] = []
    for field, column in columns.items():
        if not field.startswith("annotation__") or not isinstance(column, str):
            continue
        _, target, annotation_type, _ = field.split("__", 3)
        mappings.append(
            (column_index_from_string(column), column, target, annotation_type)
        )

    proposals: list[dict] = []
    dropped_reasons: dict[str, int] = {}
    detected_count = 0
    for _, column, target, annotation_type in sorted(mappings):
        text = _cell_text(cells, column)
        if text is None:
            continue
        if is_workflow_noise(text):
            dropped_reasons["workflow_noise"] = dropped_reasons.get("workflow_noise", 0) + 1
            continue
        detected_count += 1
        if len(proposals) >= 20:
            dropped_reasons["annotation_limit"] = dropped_reasons.get("annotation_limit", 0) + 1
            continue
        coordinate = f"{column}{row_number}"
        proposals.append(
            {
                "proposal_id": f"xlsx:{worksheet_name}:{coordinate}",
                "text": text,
                "annotation_type": annotation_type,
                "target": target,
                "source_excerpt": text,
                "source_locator": {
                    "kind": "xlsx_cell",
                    "worksheet": worksheet_name,
                    "row": row_number,
                    "column": column,
                    "coordinate": coordinate,
                },
                "provenance": "source_field",
            }
        )
    result = {"annotation_proposals": proposals}
    if dropped_reasons:
        result["dropped_annotation_count"] = sum(dropped_reasons.values())
        result["dropped_annotation_reasons"] = dropped_reasons
    omitted_count = dropped_reasons.get("annotation_limit", 0)
    if omitted_count:
        result["annotation_omitted_count"] = omitted_count
        result["annotation_detected_count"] = detected_count
    return result


def _cells_by_row_and_column(artifact: dict) -> dict[int, dict[int, dict]]:
    result: dict[int, dict[int, dict]] = {}
    for cell in artifact["cells"]:
        result.setdefault(cell["row"], {})[cell["column"]] = cell
    return result


def _header_map(
    cells_by_row: dict[int, dict[int, dict]], header_rows: list[int]
) -> dict[str, str]:
    result: dict[str, str] = {}
    for row_number in header_rows:
        for column, cell in cells_by_row.get(row_number, {}).items():
            text = _normalize_header(cell.get("displayed_text"))
            if text:
                result.setdefault(text, get_column_letter(column))
    return result


def _normalize_header(value: object) -> str:
    return " ".join(str(value or "").casefold().split())


def _cell_text(cells: dict[int, dict], column: str | None) -> str | None:
    if not column:
        return None
    cell = cells.get(column_index_from_string(column))
    if cell is None:
        return None
    value = cell.get("normalized_value")
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _line_type(value: str | None) -> Literal["material", "service", "bundled"] | None:
    normalized = " ".join((value or "").casefold().split())
    if "bundl" in normalized:
        return "bundled"
    if "material" in normalized:
        return "material"
    if "service" in normalized:
        return "service"
    return None


def _linked_concept(
    *, concept_type: Literal["material", "service"], name: str | None, category_path: str | None
) -> dict | None:
    category_suggestion = _category_suggestion(category_path)
    if name is None or category_suggestion is None:
        return None
    return {
        "concept_type": concept_type,
        "name": name,
        "category_suggestion": category_suggestion,
    }


def _category_suggestion(path: str | None) -> dict | None:
    if path is None:
        return None
    parts = [part.strip() for part in path.split("/") if part.strip()]
    if len(parts) < 2:
        return None
    return {
        "top_level_category": parts[0],
        "subcategory": " / ".join(parts[1:]),
    }


def _provider_state(value: str | None, provider_name: str | None) -> str:
    normalized = " ".join((value or "").casefold().split())
    if normalized in {"external", "internal", "unknown"}:
        return normalized
    return "external" if provider_name is not None else "unknown"
