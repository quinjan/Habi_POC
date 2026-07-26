from __future__ import annotations

import json
import shutil
from pathlib import Path

from openpyxl import load_workbook

from backend.app.xlsx.prototypes.issue_53_intelligent_xlsx.contract import (
    CellSnapshot,
    WorkbookSnapshot,
    WorksheetSnapshot,
    canonical_cell_text,
)


FIXTURE_FILENAME = "habi_issue_53_varied_completed_project.xlsx"
CONTEXT_FILENAME = "project_context.json"
CHECKED_IN_FIXTURE = (
    next(
        parent
        for parent in Path(__file__).resolve().parents
        if (parent / "AGENTS.md").is_file()
    )
    / "outputs"
    / "issue-53-intelligent-xlsx-prototype"
    / FIXTURE_FILENAME
)

EXPECTED_CANDIDATES = (
    {
        "shape": "bundle",
        "concepts": [
            {
                "kind": "material",
                "normalized_name": "Daikin split-type air conditioner",
                "observed_name_text": "Daikin split-type air-conditioning package",
                "category_path": ["Electrical", "Air Conditioning Equipment"],
                "quantity": None,
                "unit": None,
                "component_unit_price": None,
            },
            {
                "kind": "service",
                "normalized_name": "Air conditioning installation",
                "observed_name_text": "installation",
                "category_path": ["Services", "Air Conditioning Installation"],
                "quantity": None,
                "unit": None,
                "component_unit_price": None,
            },
        ],
        "provider_state": "external",
        "provider_name": "CoolAir Mechanical Services",
        "observed_provider_text": "CoolAir Mechanical Services",
        "provider_roles": [
            "material_supplier",
            "service_provider",
            "supply_and_install_provider",
        ],
        "provider_category_path": ["Providers", "HVAC Contractors"],
        "commercial_quantity": "1",
        "commercial_unit": "package",
        "currency": "PHP",
        "unit_price": None,
        "total_price": "120000",
    },
    {
        "shape": "material",
        "concepts": [
            {
                "kind": "material",
                "normalized_name": "100 mm PVC pressure pipe",
                "observed_name_text": "100 mm PVC pressure pipes",
                "category_path": ["Plumbing", "Pipes"],
                "quantity": "20",
                "unit": "lengths",
                "component_unit_price": "1625",
            }
        ],
        "provider_state": "external",
        "provider_name": "BuildMart Trading",
        "observed_provider_text": "BuildMart Trading",
        "provider_roles": ["material_supplier"],
        "provider_category_path": [
            "Providers",
            "Construction Material Suppliers",
        ],
        "commercial_quantity": "20",
        "commercial_unit": "lengths",
        "currency": "PHP",
        "unit_price": "1625",
        "total_price": "32500",
    },
    {
        "shape": "material",
        "concepts": [
            {
                "kind": "material",
                "normalized_name": "16 mm Grade 60 deformed reinforcing bar",
                "observed_name_text": "16 mm Grade 60 deformed reinforcing bars",
                "category_path": ["Civil", "Reinforcing Steel"],
                "quantity": "24",
                "unit": "lengths",
                "component_unit_price": "1200",
            }
        ],
        "provider_state": "external",
        "provider_name": "MetroSteel",
        "observed_provider_text": "MetroSteel",
        "provider_roles": ["material_supplier"],
        "provider_category_path": [
            "Providers",
            "Construction Material Suppliers",
        ],
        "commercial_quantity": "24",
        "commercial_unit": "lengths",
        "currency": "PHP",
        "unit_price": "1200",
        "total_price": "28800",
    },
    {
        "shape": "service",
        "concepts": [
            {
                "kind": "service",
                "normalized_name": "Site cleanup",
                "observed_name_text": "Final site cleanup",
                "category_path": ["Services", "Site Cleanup"],
                "quantity": "1",
                "unit": "job",
                "component_unit_price": None,
            }
        ],
        "provider_state": "internal",
        "provider_name": None,
        "observed_provider_text": "Our own crew",
        "provider_roles": ["service_provider"],
        "provider_category_path": None,
        "commercial_quantity": "1",
        "commercial_unit": "job",
        "currency": "PHP",
        "unit_price": None,
        "total_price": "18000",
    },
    {
        "shape": "material",
        "concepts": [
            {
                "kind": "material",
                "normalized_name": "Non-shrink grout",
                "observed_name_text": "Non-shrink grout",
                "category_path": ["Civil", "Grouting Materials"],
                "quantity": "10",
                "unit": "bags",
                "component_unit_price": "850",
            }
        ],
        "provider_state": "unknown",
        "provider_name": None,
        "observed_provider_text": None,
        "provider_roles": [],
        "provider_category_path": None,
        "commercial_quantity": "10",
        "commercial_unit": "bags",
        "currency": "PHP",
        "unit_price": "850",
        "total_price": "8500",
    },
    {
        "shape": "service",
        "concepts": [
            {
                "kind": "service",
                "normalized_name": "Domestic water lines pressure testing",
                "observed_name_text": (
                    "Final pressure testing of domestic water lines"
                ),
                "category_path": [
                    "Services",
                    "Testing and Commissioning",
                ],
                "quantity": "1",
                "unit": "job",
                "component_unit_price": None,
            }
        ],
        "provider_state": "internal",
        "provider_name": None,
        "observed_provider_text": "Arnaiz Builders crew",
        "provider_roles": ["service_provider"],
        "provider_category_path": None,
        "commercial_quantity": "1",
        "commercial_unit": "job",
        "currency": "PHP",
        "unit_price": None,
        "total_price": "9000",
    },
)

REQUIRED_OMISSIONS = (
    "Ceiling installation",
    "preliminary quote",
    "For approval",
    "Follow up with Mario",
    "TBD",
)


def create_fixture(output_dir: Path) -> tuple[Path, Path]:
    if not CHECKED_IN_FIXTURE.is_file():
        raise RuntimeError(f"Prototype fixture is missing: {CHECKED_IN_FIXTURE}")
    output_dir.mkdir(parents=True, exist_ok=True)
    workbook_path = output_dir / FIXTURE_FILENAME
    context_path = output_dir / CONTEXT_FILENAME
    if workbook_path.resolve() != CHECKED_IN_FIXTURE.resolve():
        shutil.copyfile(CHECKED_IN_FIXTURE, workbook_path)
    context_path.write_text(
        json.dumps(_project_context(), indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    return workbook_path, context_path


def load_workbook_snapshot(workbook_path: Path) -> WorkbookSnapshot:
    formula_workbook = load_workbook(
        workbook_path,
        data_only=False,
        read_only=False,
    )
    value_workbook = load_workbook(
        workbook_path,
        data_only=True,
        read_only=False,
    )
    try:
        worksheets: dict[str, WorksheetSnapshot] = {}
        for formula_sheet in formula_workbook.worksheets:
            value_sheet = value_workbook[formula_sheet.title]
            cells: dict[str, CellSnapshot] = {}
            for row in formula_sheet.iter_rows():
                for formula_cell in row:
                    value_cell = value_sheet[formula_cell.coordinate]
                    if formula_cell.value is None and value_cell.value is None:
                        continue
                    cells[formula_cell.coordinate] = CellSnapshot(
                        raw_text=canonical_cell_text(formula_cell.value),
                        cached_text=canonical_cell_text(value_cell.value),
                    )
            worksheets[formula_sheet.title] = WorksheetSnapshot(
                visibility=formula_sheet.sheet_state,
                used_range=formula_sheet.calculate_dimension() if cells else None,
                max_row=formula_sheet.max_row,
                max_column=formula_sheet.max_column,
                cells=cells,
            )
    finally:
        formula_workbook.close()
        value_workbook.close()
    return WorkbookSnapshot(
        filename=workbook_path.name,
        worksheets=worksheets,
    )


def _project_context() -> dict:
    return {
        "project_workspace": {
            "project_name": "Arnaiz Residence",
            "contractor_assigned": "Arnaiz Builders",
            "currency_default": "PHP",
        },
        "project_memory": [
            {
                "record_id": 101,
                "record_type": "material",
                "name": "Daikin split-type air conditioner",
                "category_path": [
                    "Electrical",
                    "Air Conditioning Equipment",
                ],
            },
            {
                "record_id": 102,
                "record_type": "provider",
                "name": "CoolAir Mechanical Services",
                "category_path": ["Providers", "HVAC Contractors"],
            },
        ],
        "taxonomy_vocabulary": [
            ["Electrical", "Air Conditioning Equipment"],
            ["Services", "Air Conditioning Installation"],
            ["Providers", "HVAC Contractors"],
            ["Plumbing", "Pipes"],
            ["Providers", "Construction Material Suppliers"],
            ["Services", "Site Cleanup"],
            ["Civil", "Grouting Materials"],
            ["Civil", "Reinforcing Steel"],
            ["Services", "Testing and Commissioning"],
        ],
        "prototype_rules": {
            "source_file_is_final_as_used_by_default": True,
            "candidate_local_non_final_wording_overrides_default": True,
            "bid_and_proposal_workbooks_are_ineligible": True,
            "every_candidate_requires_source_evidence": True,
            "reviewer_approval_is_required_before_import": True,
        },
    }
