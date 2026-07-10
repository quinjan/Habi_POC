from datetime import date, datetime, time
from decimal import Decimal
import json
import os
from pathlib import Path
import tempfile

from openpyxl import Workbook
from openpyxl.utils import column_index_from_string
from sqlalchemy.orm import Session

from backend.app.sources.models import SourceFile
from backend.app.xlsx.config import XlsxProcessingConfig
from backend.app.xlsx.models import WorksheetArtifact


class WorkbookLimitExceeded(ValueError):
    def __init__(self, message: str, diagnostics: dict):
        super().__init__(message)
        self.diagnostics = diagnostics


def create_worksheet_artifacts(
    *,
    session: Session,
    source_file: SourceFile,
    formula_workbook: Workbook,
    value_workbook: Workbook,
    config: XlsxProcessingConfig,
) -> tuple[list[WorksheetArtifact], dict]:
    payloads, diagnostics = _scan_visible_worksheets(
        formula_workbook, value_workbook, config
    )

    artifacts: list[WorksheetArtifact] = []
    for payload in payloads:
        worksheet = payload["worksheet"]
        relative_path = (
            Path("derived")
            / str(source_file.id)
            / "worksheets"
            / f"{worksheet['index']}.json"
        )
        _write_json_atomically(config.storage_root / relative_path, payload)
        artifact = WorksheetArtifact(
            project_workspace_id=source_file.project_workspace_id,
            source_file_id=source_file.id,
            worksheet_index=worksheet["index"],
            worksheet_name=worksheet["name"],
            artifact_path=relative_path.as_posix(),
            non_empty_row_count=payload["non_empty_row_count"],
            non_empty_cell_count=payload["non_empty_cell_count"],
            profile=None,
        )
        session.add(artifact)
        session.flush()
        artifacts.append(artifact)

    diagnostics["artifact_count"] = len(artifacts)
    return artifacts, diagnostics


def _scan_visible_worksheets(
    formula_workbook: Workbook,
    value_workbook: Workbook,
    config: XlsxProcessingConfig,
) -> tuple[list[dict], dict]:
    payloads: list[dict] = []
    visible_worksheet_count = sum(
        sheet.sheet_state == "visible" for sheet in formula_workbook.worksheets
    )
    hidden_worksheet_count = len(formula_workbook.worksheets) - visible_worksheet_count
    hidden_row_count = 0
    hidden_column_count = 0
    total_non_empty_rows = 0
    total_non_empty_cells = 0
    _raise_if_limit_exceeded(
        label="visible worksheets",
        actual=visible_worksheet_count,
        maximum=config.max_visible_worksheets,
        diagnostics=_scan_diagnostics(
            visible_worksheet_count=visible_worksheet_count,
            hidden_worksheet_count=hidden_worksheet_count,
            hidden_row_count=0,
            hidden_column_count=0,
            non_empty_row_count=0,
            non_empty_cell_count=0,
        ),
    )

    for worksheet_index, formula_sheet in enumerate(formula_workbook.worksheets):
        if formula_sheet.sheet_state != "visible":
            continue
        value_sheet = value_workbook[formula_sheet.title]
        hidden_rows = {
            index
            for index, dimension in formula_sheet.row_dimensions.items()
            if dimension.hidden and index <= formula_sheet.max_row
        }
        hidden_columns: set[int] = set()
        for key, dimension in formula_sheet.column_dimensions.items():
            if not dimension.hidden:
                continue
            start = dimension.min or column_index_from_string(key)
            end = dimension.max or start
            hidden_columns.update(
                column_index
                for column_index in range(start, end + 1)
                if column_index <= formula_sheet.max_column
            )
        hidden_row_count += len(hidden_rows)
        hidden_column_count += len(hidden_columns)

        cells: list[dict] = []
        non_empty_rows: set[int] = set()
        for row in formula_sheet.iter_rows():
            for formula_cell in row:
                if formula_cell.row in hidden_rows or formula_cell.column in hidden_columns:
                    continue
                value_cell = value_sheet[formula_cell.coordinate]
                formula = formula_cell.value if formula_cell.data_type == "f" else None
                raw_value = value_cell.value if formula is not None else formula_cell.value
                if formula_cell.value is None and raw_value is None:
                    continue
                if formula_cell.row not in non_empty_rows:
                    non_empty_rows.add(formula_cell.row)
                    total_non_empty_rows += 1
                    _raise_if_limit_exceeded(
                        label="visible non-empty rows",
                        actual=total_non_empty_rows,
                        maximum=config.max_non_empty_rows,
                        diagnostics=_scan_diagnostics(
                            visible_worksheet_count=visible_worksheet_count,
                            hidden_worksheet_count=hidden_worksheet_count,
                            hidden_row_count=hidden_row_count,
                            hidden_column_count=hidden_column_count,
                            non_empty_row_count=total_non_empty_rows,
                            non_empty_cell_count=total_non_empty_cells,
                        ),
                    )
                cells.append(
                    _cell_payload(
                        formula_cell=formula_cell,
                        raw_value=raw_value,
                        formula=formula,
                    )
                )
                total_non_empty_cells += 1
                _raise_if_limit_exceeded(
                    label="visible non-empty cells",
                    actual=total_non_empty_cells,
                    maximum=config.max_non_empty_cells,
                    diagnostics=_scan_diagnostics(
                        visible_worksheet_count=visible_worksheet_count,
                        hidden_worksheet_count=hidden_worksheet_count,
                        hidden_row_count=hidden_row_count,
                        hidden_column_count=hidden_column_count,
                        non_empty_row_count=total_non_empty_rows,
                        non_empty_cell_count=total_non_empty_cells,
                    ),
                )

        payloads.append(
            {
                "worksheet": {
                    "name": formula_sheet.title,
                    "index": worksheet_index,
                    "merged_ranges": [
                        str(cell_range) for cell_range in formula_sheet.merged_cells.ranges
                    ],
                },
                "cells": cells,
                "non_empty_row_count": len(non_empty_rows),
                "non_empty_cell_count": len(cells),
            }
        )
    diagnostics = _scan_diagnostics(
        visible_worksheet_count=visible_worksheet_count,
        hidden_worksheet_count=hidden_worksheet_count,
        hidden_row_count=hidden_row_count,
        hidden_column_count=hidden_column_count,
        non_empty_row_count=total_non_empty_rows,
        non_empty_cell_count=total_non_empty_cells,
    )
    return payloads, diagnostics


def _cell_payload(*, formula_cell, raw_value: object, formula: str | None) -> dict:
    normalized_value = _json_value(raw_value)
    return {
        "coordinate": formula_cell.coordinate,
        "row": formula_cell.row,
        "column": formula_cell.column,
        "raw_value": normalized_value,
        "normalized_value": normalized_value,
        "displayed_text": _displayed_text(normalized_value),
        "cell_type": _cell_type(formula_cell, raw_value, formula),
        "number_format": formula_cell.number_format,
        "formula": formula,
        "is_bold": bool(formula_cell.font and formula_cell.font.bold),
        "has_fill": bool(formula_cell.fill and formula_cell.fill.fill_type),
    }


def _json_value(value: object) -> object:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _displayed_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def _cell_type(formula_cell, raw_value: object, formula: str | None) -> str:
    if formula is not None:
        return "formula"
    if formula_cell.is_date or isinstance(raw_value, (date, datetime, time)):
        return "date"
    if isinstance(raw_value, bool):
        return "boolean"
    if isinstance(raw_value, (int, float, Decimal)):
        return "number"
    if formula_cell.data_type == "e":
        return "error"
    return "string"


def _scan_diagnostics(
    *,
    visible_worksheet_count: int,
    hidden_worksheet_count: int,
    hidden_row_count: int,
    hidden_column_count: int,
    non_empty_row_count: int,
    non_empty_cell_count: int,
) -> dict:
    return {
        "visible_worksheet_count": visible_worksheet_count,
        "hidden_worksheet_count": hidden_worksheet_count,
        "hidden_row_count": hidden_row_count,
        "hidden_column_count": hidden_column_count,
        "visible_non_empty_row_count": non_empty_row_count,
        "visible_non_empty_cell_count": non_empty_cell_count,
    }


def _raise_if_limit_exceeded(
    *, label: str, actual: int, maximum: int, diagnostics: dict
) -> None:
    if actual > maximum:
        raise WorkbookLimitExceeded(
            f"Workbook exceeds the configured {label} limit ({actual} > {maximum})",
            diagnostics,
        )


def _write_json_atomically(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            dir=path.parent,
            delete=False,
            suffix=".json.tmp",
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
            json.dump(payload, temporary_file, ensure_ascii=False, separators=(",", ":"))
        os.replace(temporary_path, path)
        temporary_path = None
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
