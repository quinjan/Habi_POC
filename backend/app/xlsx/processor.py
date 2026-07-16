import json
import inspect
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from backend.app.processing.ai_extraction import apply_contractor_assigned_provider_default
from backend.app.processing.memory_context import build_project_memory_context
from backend.app.processing.models import ProcessingJob
from backend.app.review.models import ExtractedCandidate, ReviewBatch
from backend.app.sources.models import SourceFile
from backend.app.xlsx.artifacts import WorkbookLimitExceeded, create_worksheet_artifacts
from backend.app.xlsx.ai import validate_worksheet_profile, validate_xlsx_candidates
from backend.app.xlsx.config import XlsxProcessingConfig
from backend.app.xlsx.grounding import (
    build_explicit_row_candidates,
    ground_profile_to_source,
    sanitize_profile_column_mappings,
)


def process_xlsx_source_file(
    session: Session,
    job: ProcessingJob,
    ai_provider: object,
    config: XlsxProcessingConfig,
) -> tuple[ReviewBatch | None, list[ExtractedCandidate], dict, str, str | None]:
    provider_name = getattr(ai_provider, "provider_name", "fake")
    provider_model = getattr(ai_provider, "model", "fake-ai-provider")
    diagnostics = {
        "processor": "ai_xlsx_purchase_lines_v1",
        "provider": provider_name,
        "model": provider_model,
    }
    source_file = session.scalar(
        select(SourceFile).where(
            SourceFile.source_submission_id == job.source_submission_id,
            SourceFile.project_workspace_id == job.project_workspace_id,
        )
    )
    if source_file is None:
        message = "XLSX Processing Job requires a preserved Source File"
        diagnostics["failure_summary"] = message
        return None, [], diagnostics, "failed", message

    workbook_path = config.storage_root / Path(source_file.storage_path)
    diagnostics["original_filename"] = source_file.original_filename
    try:
        formula_workbook = load_workbook(workbook_path, data_only=False, read_only=False)
        value_workbook = load_workbook(workbook_path, data_only=True, read_only=False)
    except Exception as error:
        message = f"Workbook could not be opened: {error}"
        diagnostics.update(
            failure="workbook_open_error",
            failure_summary=message,
        )
        return None, [], diagnostics, "failed", message

    try:
        artifacts, artifact_diagnostics = create_worksheet_artifacts(
            session=session,
            source_file=source_file,
            formula_workbook=formula_workbook,
            value_workbook=value_workbook,
            config=config,
        )
        diagnostics.update(artifact_diagnostics)
    except WorkbookLimitExceeded as error:
        message = str(error)
        diagnostics.update(error.diagnostics)
        diagnostics.update(failure="workbook_limit_exceeded", failure_summary=message)
        return None, [], diagnostics, "failed", message
    except Exception as error:
        message = f"Worksheet artifacts could not be created: {error}"
        diagnostics.update(failure="artifact_creation_error", failure_summary=message)
        return None, [], diagnostics, "failed", message
    finally:
        formula_workbook.close()
        value_workbook.close()

    valid_payloads: list[dict] = []
    raw_candidate_count = 0
    dropped_candidate_count = 0
    profile_request_count = 0
    extraction_request_count = 0
    extraction_chunk_count = 0
    usable_region_count = 0
    unusable_region_count = 0
    source_grounded_candidate_count = 0
    ai_candidate_replaced_count = 0
    invalid_profile_column_mapping_count = 0
    artifact_contents = [
        json.loads((config.storage_root / artifact.artifact_path).read_text(encoding="utf-8"))
        for artifact in artifacts
    ]
    source_text = " ".join(
        str(cell.get("displayed_text") or cell.get("raw_value") or "")
        for artifact_content in artifact_contents
        for cell in artifact_content["cells"]
    )
    memory_context, omitted_counts = build_project_memory_context(
        session=session,
        project_workspace_id=job.project_workspace_id,
        source_text=source_text,
    )
    try:
        for artifact, artifact_content in zip(artifacts, artifact_contents, strict=True):
            profile_request_count += 1
            raw_profile = ai_provider.profile_worksheet(
                worksheet=artifact_content,
                source_submission_id=job.source_submission_id,
            )
            raw_profile, invalid_mapping_count = sanitize_profile_column_mappings(
                raw_profile, artifact_content
            )
            invalid_profile_column_mapping_count += invalid_mapping_count
            profile = validate_worksheet_profile(raw_profile, artifact_content)
            profile = ground_profile_to_source(profile, artifact_content)
            profile = validate_worksheet_profile(profile, artifact_content)
            artifact.profile = profile

            row_payloads = _artifact_rows_for_ai(artifact_content)
            rows_by_number = {row["row"]: row for row in row_payloads}
            for region in profile["regions"]:
                if not region["usable"]:
                    unusable_region_count += 1
                    continue
                usable_region_count += 1
                body_rows = [
                    row
                    for row in row_payloads
                    if region["body_start_row"] <= row["row"] <= region["body_end_row"]
                ]
                context_row_numbers = list(
                    dict.fromkeys(
                        profile["title_rows"]
                        + profile["header_rows"]
                        + region["header_row_numbers"]
                    )
                )
                context_rows = [
                    rows_by_number[row_number]
                    for row_number in context_row_numbers
                    if row_number in rows_by_number
                    and not (
                        region["body_start_row"]
                        <= row_number
                        <= region["body_end_row"]
                    )
                ]
                for chunk in _chunks(body_rows, config.extraction_chunk_rows):
                    extraction_chunk_count += 1
                    extraction_request_count += 1
                    provider_region = {
                        **region,
                        "source_file_id": source_file.id,
                        "worksheet": artifact_content["worksheet"]["name"],
                    }
                    extract = ai_provider.extract_worksheet_chunk
                    kwargs = {
                        "profile": profile,
                        "region": provider_region,
                        "rows": chunk,
                        "context_rows": context_rows,
                        "source_submission_id": job.source_submission_id,
                    }
                    if "memory_context" in inspect.signature(extract).parameters:
                        kwargs["memory_context"] = memory_context
                    raw_result = extract(**kwargs)
                    if not isinstance(raw_result, dict) or not isinstance(
                        raw_result.get("candidates", []), list
                    ):
                        raise ValueError("AI provider returned malformed XLSX candidates")
                    raw_candidates = raw_result.get("candidates", [])
                    raw_candidate_count += len(raw_candidates)
                    valid, dropped = validate_xlsx_candidates(
                        raw_candidates=raw_candidates,
                        source_submission_id=job.source_submission_id,
                        source_file_id=source_file.id,
                        artifact=artifact_content,
                        profile=profile,
                        expected_region_id=region["region_id"],
                    )
                    explicit_candidates = build_explicit_row_candidates(
                        rows=chunk,
                        region=region,
                        source_submission_id=job.source_submission_id,
                        source_file_id=source_file.id,
                        worksheet_name=artifact_content["worksheet"]["name"],
                    )
                    explicit_row_numbers = set(explicit_candidates)
                    ai_candidate_replaced_count += sum(
                        payload["evidence"]["primary_body_row"] in explicit_row_numbers
                        for payload in valid
                    )
                    ungrounded_ai_candidates = [
                        payload
                        for payload in valid
                        if payload["evidence"]["primary_body_row"]
                        not in explicit_row_numbers
                    ]
                    grounded_valid, grounded_dropped = validate_xlsx_candidates(
                        raw_candidates=list(explicit_candidates.values()),
                        source_submission_id=job.source_submission_id,
                        source_file_id=source_file.id,
                        artifact=artifact_content,
                        profile=profile,
                        expected_region_id=region["region_id"],
                        ground_ai_annotations=False,
                    )
                    source_grounded_candidate_count += len(grounded_valid)
                    chunk_payloads = ungrounded_ai_candidates + grounded_valid
                    chunk_payloads.sort(
                        key=lambda payload: payload["evidence"]["primary_body_row"]
                    )
                    valid_payloads.extend(
                        apply_contractor_assigned_provider_default(
                            payload,
                            contractor_assigned=memory_context["contractor_assigned"],
                        )
                        for payload in chunk_payloads
                    )
                    dropped_candidate_count += dropped + grounded_dropped
    except Exception as error:
        message = f"XLSX AI processing failed: {error}"
        diagnostics.update(
            failure="ai_processing_error",
            failure_summary=message,
            profile_request_count=profile_request_count,
            extraction_request_count=extraction_request_count,
            invalid_profile_column_mapping_count=invalid_profile_column_mapping_count,
        )
        return None, [], diagnostics, "failed", message

    diagnostics.update(
        profile_request_count=profile_request_count,
        extraction_request_count=extraction_request_count,
        extraction_chunk_count=extraction_chunk_count,
        configured_chunk_row_limit=config.extraction_chunk_rows,
        usable_region_count=usable_region_count,
        unusable_region_count=unusable_region_count,
        raw_candidate_count=raw_candidate_count,
        valid_candidate_count=len(valid_payloads),
        dropped_candidate_count=dropped_candidate_count,
        source_grounded_candidate_count=source_grounded_candidate_count,
        ai_candidate_replaced_count=ai_candidate_replaced_count,
        invalid_profile_column_mapping_count=invalid_profile_column_mapping_count,
    )
    if any(omitted_counts.values()):
        diagnostics["memory_context_omitted_counts"] = omitted_counts
    if not valid_payloads:
        return None, [], diagnostics, "no_candidates_found", None

    review_batch = ReviewBatch(
        project_workspace_id=job.project_workspace_id,
        source_submission_id=job.source_submission_id,
        status="review_pending",
    )
    session.add(review_batch)
    session.flush()
    candidates: list[ExtractedCandidate] = []
    for payload in valid_payloads:
        candidate = ExtractedCandidate(
            project_workspace_id=job.project_workspace_id,
            review_batch_id=review_batch.id,
            source_submission_id=job.source_submission_id,
            status="pending_review",
            proposed_payload=payload,
        )
        session.add(candidate)
        candidates.append(candidate)
    session.flush()
    return review_batch, candidates, diagnostics, "review_ready", None


def _artifact_rows_for_ai(artifact: dict) -> list[dict]:
    rows: dict[int, list[dict]] = {}
    for cell in artifact["cells"]:
        if cell["cell_type"] == "formula" and cell["raw_value"] is None:
            continue
        rows.setdefault(cell["row"], []).append(cell)
    return [
        {"row": row_number, "cells": cells}
        for row_number, cells in sorted(rows.items())
    ]


def _chunks(rows: list[dict], size: int):
    for index in range(0, len(rows), size):
        yield rows[index : index + size]
