from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries


SUBMISSION_VERSION = "issue-53-prototype-v1"
DISPOSITIONS = (
    "context",
    "header",
    "candidate_data",
    "supporting_context",
    "excluded_non_final",
    "noise",
)


def submit_candidate_batch_tool() -> dict[str, Any]:
    return {
        "type": "function",
        "name": "submit_candidate_batch",
        "description": (
            "Submit the one complete provisional Habi candidate batch. Call exactly "
            "once, after inspecting the entire workbook. This is the only accepted "
            "Habi output boundary. Every important candidate field must cite one or "
            "more independently verifiable evidence ids."
        ),
        "strict": True,
        "parameters": _submission_schema(),
    }


def _nullable_string() -> dict[str, Any]:
    return {"type": ["string", "null"]}


def _string_array() -> dict[str, Any]:
    return {"type": "array", "items": {"type": "string"}}


def _submission_schema() -> dict[str, Any]:
    range_accounting = {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "range": {"type": "string"},
            "disposition": {"type": "string", "enum": list(DISPOSITIONS)},
            "candidate_ids": _string_array(),
            "rationale": {"type": "string"},
        },
        "required": ["range", "disposition", "candidate_ids", "rationale"],
    }
    worksheet = {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "name": {"type": "string"},
            "visibility": {
                "type": "string",
                "enum": ["visible", "hidden", "veryHidden"],
            },
            "used_range": _nullable_string(),
            "accounting": {"type": "array", "items": range_accounting},
        },
        "required": ["name", "visibility", "used_range", "accounting"],
    }
    evidence = {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "evidence_id": {"type": "string"},
            "worksheet": {"type": "string"},
            "range": {"type": "string"},
            "quoted_cells": {
                "type": "array",
                "minItems": 1,
                "items": {
                    "type": "object",
                    "additionalProperties": False,
                    "properties": {
                        "coordinate": {"type": "string"},
                        "text": {"type": "string"},
                    },
                    "required": ["coordinate", "text"],
                },
            },
            "purpose": {"type": "string"},
        },
        "required": [
            "evidence_id",
            "worksheet",
            "range",
            "quoted_cells",
            "purpose",
        ],
    }
    concept = {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "kind": {"type": "string", "enum": ["material", "service"]},
            "normalized_name": {"type": "string"},
            "observed_name_text": {"type": "string"},
            "category_path": {
                "type": "array",
                "minItems": 2,
                "maxItems": 2,
                "items": {"type": "string"},
            },
            "quantity": _nullable_string(),
            "unit": _nullable_string(),
            "component_unit_price": _nullable_string(),
        },
        "required": [
            "kind",
            "normalized_name",
            "observed_name_text",
            "category_path",
            "quantity",
            "unit",
            "component_unit_price",
        ],
    }
    provider = {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "state": {
                "type": "string",
                "enum": ["external", "internal", "unknown"],
            },
            "name": _nullable_string(),
            "observed_provider_text": _nullable_string(),
            "roles": {
                "type": "array",
                "items": {
                    "type": "string",
                    "enum": [
                        "material_supplier",
                        "service_provider",
                        "supply_and_install_provider",
                    ],
                },
            },
            "category_path": {
                "type": ["array", "null"],
                "items": {"type": "string"},
                "minItems": 2,
                "maxItems": 2,
            },
        },
        "required": [
            "state",
            "name",
            "observed_provider_text",
            "roles",
            "category_path",
        ],
    }
    candidate = {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "candidate_id": {"type": "string"},
            "shape": {
                "type": "string",
                "enum": ["material", "service", "bundle"],
            },
            "purchasing_status": {"type": "string", "enum": ["final"]},
            "concepts": {"type": "array", "minItems": 1, "items": concept},
            "provider": provider,
            "commercial_quantity": _nullable_string(),
            "commercial_unit": _nullable_string(),
            "currency": _nullable_string(),
            "unit_price": _nullable_string(),
            "total_price": _nullable_string(),
            "purchase_date": _nullable_string(),
            "field_evidence": {
                "type": "array",
                "minItems": 1,
                "items": {
                    "type": "object",
                    "additionalProperties": False,
                    "properties": {
                        "field_path": {"type": "string"},
                        "evidence_ids": {
                            "type": "array",
                            "minItems": 1,
                            "items": {"type": "string"},
                        },
                    },
                    "required": ["field_path", "evidence_ids"],
                },
            },
        },
        "required": [
            "candidate_id",
            "shape",
            "purchasing_status",
            "concepts",
            "provider",
            "commercial_quantity",
            "commercial_unit",
            "currency",
            "unit_price",
            "total_price",
            "purchase_date",
            "field_evidence",
        ],
    }
    exclusion = {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "exclusion_id": {"type": "string"},
            "reason": {
                "type": "string",
                "enum": [
                    "non_final",
                    "workflow_noise",
                    "unsupported",
                    "ineligible",
                ],
            },
            "description": {"type": "string"},
            "evidence_ids": {
                "type": "array",
                "minItems": 1,
                "items": {"type": "string"},
            },
        },
        "required": ["exclusion_id", "reason", "description", "evidence_ids"],
    }
    return {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "submission_version": {
                "type": "string",
                "enum": [SUBMISSION_VERSION],
            },
            "workbook_filename": {"type": "string"},
            "worksheets": {"type": "array", "minItems": 1, "items": worksheet},
            "candidates": {"type": "array", "items": candidate},
            "evidence": {"type": "array", "items": evidence},
            "exclusions": {"type": "array", "items": exclusion},
        },
        "required": [
            "submission_version",
            "workbook_filename",
            "worksheets",
            "candidates",
            "evidence",
            "exclusions",
        ],
    }


@dataclass(frozen=True)
class VerificationResult:
    passed: bool
    errors: tuple[str, ...]
    worksheet_count: int
    visible_worksheet_count: int
    non_empty_cell_count: int
    accounted_non_empty_cell_count: int
    candidate_count: int
    evidence_count: int


def verify_submission(
    workbook_path: Path,
    submission: dict[str, Any],
) -> VerificationResult:
    errors: list[str] = []
    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    try:
        inventory = {
            sheet.title: {
                "visibility": sheet.sheet_state,
                "used_range": _used_range(sheet),
                "non_empty": _non_empty_coordinates(sheet),
            }
            for sheet in workbook.worksheets
        }
        worksheets = _list(submission, "worksheets", errors)
        submitted_names = [item.get("name") for item in worksheets if isinstance(item, dict)]
        if len(submitted_names) != len(set(submitted_names)):
            errors.append("Each worksheet must be inventoried exactly once")
        if set(submitted_names) != set(inventory):
            errors.append(
                "Worksheet inventory mismatch: "
                f"expected {sorted(inventory)}, received {sorted(str(x) for x in submitted_names)}"
            )

        candidate_list = _list(submission, "candidates", errors)
        candidates = {
            item.get("candidate_id"): item
            for item in candidate_list
            if isinstance(item, dict) and isinstance(item.get("candidate_id"), str)
        }
        if len(candidates) != len(candidate_list):
            errors.append("Candidate ids must be present and unique")

        evidence_list = _list(submission, "evidence", errors)
        evidence = {
            item.get("evidence_id"): item
            for item in evidence_list
            if isinstance(item, dict) and isinstance(item.get("evidence_id"), str)
        }
        if len(evidence) != len(evidence_list):
            errors.append("Evidence ids must be present and unique")

        accounted_non_empty: set[tuple[str, str]] = set()
        for item in worksheets:
            if not isinstance(item, dict):
                errors.append("Worksheet accounting entries must be objects")
                continue
            name = item.get("name")
            if name not in inventory:
                continue
            expected = inventory[name]
            if item.get("visibility") != expected["visibility"]:
                errors.append(f"{name}: visibility does not match workbook")
            if item.get("used_range") != expected["used_range"]:
                errors.append(
                    f"{name}: used_range expected {expected['used_range']!r}, "
                    f"received {item.get('used_range')!r}"
                )
            accounting = item.get("accounting")
            if not isinstance(accounting, list):
                errors.append(f"{name}: accounting must be an array")
                continue
            if expected["visibility"] != "visible":
                if accounting:
                    errors.append(f"{name}: hidden worksheet must not have accounting ranges")
                continue
            sheet = workbook[name]
            for region in accounting:
                if not isinstance(region, dict):
                    errors.append(f"{name}: accounting region must be an object")
                    continue
                coordinates = _coordinates_in_range(
                    sheet,
                    region.get("range"),
                    errors,
                    f"{name} accounting",
                )
                for candidate_id in region.get("candidate_ids", []):
                    if candidate_id not in candidates:
                        errors.append(
                            f"{name}: accounting references unknown candidate {candidate_id!r}"
                        )
                for coordinate in coordinates & expected["non_empty"]:
                    accounted_non_empty.add((name, coordinate))

        all_visible_non_empty = {
            (name, coordinate)
            for name, sheet_inventory in inventory.items()
            if sheet_inventory["visibility"] == "visible"
            for coordinate in sheet_inventory["non_empty"]
        }
        missing_coverage = sorted(all_visible_non_empty - accounted_non_empty)
        if missing_coverage:
            preview = ", ".join(f"{sheet}!{cell}" for sheet, cell in missing_coverage[:12])
            errors.append(
                f"{len(missing_coverage)} visible non-empty cells are unaccounted for: {preview}"
            )

        for evidence_id, item in evidence.items():
            worksheet_name = item.get("worksheet")
            if worksheet_name not in inventory:
                errors.append(f"{evidence_id}: unknown worksheet {worksheet_name!r}")
                continue
            if inventory[worksheet_name]["visibility"] != "visible":
                errors.append(f"{evidence_id}: hidden worksheet cannot supply evidence")
            sheet = workbook[worksheet_name]
            coordinates = _coordinates_in_range(
                sheet,
                item.get("range"),
                errors,
                f"Evidence {evidence_id}",
            )
            quoted_cells = item.get("quoted_cells")
            if not isinstance(quoted_cells, list) or not quoted_cells:
                errors.append(f"{evidence_id}: quoted_cells must be non-empty")
                continue
            for quote in quoted_cells:
                if not isinstance(quote, dict):
                    errors.append(f"{evidence_id}: quoted cell must be an object")
                    continue
                coordinate = str(quote.get("coordinate", "")).upper()
                if coordinate not in coordinates:
                    errors.append(
                        f"{evidence_id}: quoted cell {coordinate!r} is outside evidence range"
                    )
                    continue
                expected_text = canonical_cell_text(sheet[coordinate].value)
                if quote.get("text") != expected_text:
                    errors.append(
                        f"{evidence_id}: {worksheet_name}!{coordinate} exact text mismatch; "
                        f"expected {expected_text!r}, received {quote.get('text')!r}"
                    )

        for candidate_id, candidate in candidates.items():
            _verify_candidate(candidate_id, candidate, evidence, errors)
        for exclusion in _list(submission, "exclusions", errors):
            if not isinstance(exclusion, dict):
                errors.append("Exclusions must be objects")
                continue
            _verify_evidence_references(
                f"Exclusion {exclusion.get('exclusion_id')}",
                exclusion.get("evidence_ids"),
                evidence,
                errors,
            )
    finally:
        workbook.close()

    return VerificationResult(
        passed=not errors,
        errors=tuple(errors),
        worksheet_count=len(inventory),
        visible_worksheet_count=sum(
            1 for item in inventory.values() if item["visibility"] == "visible"
        ),
        non_empty_cell_count=sum(
            len(item["non_empty"])
            for item in inventory.values()
            if item["visibility"] == "visible"
        ),
        accounted_non_empty_cell_count=len(accounted_non_empty),
        candidate_count=len(candidates),
        evidence_count=len(evidence),
    )


def _verify_candidate(
    candidate_id: str,
    candidate: dict[str, Any],
    evidence: dict[str, dict[str, Any]],
    errors: list[str],
) -> None:
    concepts = candidate.get("concepts")
    if not isinstance(concepts, list) or not concepts:
        errors.append(f"{candidate_id}: concepts must be non-empty")
        return
    shape = candidate.get("shape")
    if shape == "bundle" and len(concepts) < 2:
        errors.append(f"{candidate_id}: bundle needs at least two concepts")
    if shape in {"material", "service"}:
        if len(concepts) != 1 or concepts[0].get("kind") != shape:
            errors.append(f"{candidate_id}: {shape} shape must have one matching concept")

    provider = candidate.get("provider")
    if not isinstance(provider, dict):
        errors.append(f"{candidate_id}: provider must be an object")
        return
    provider_state = provider.get("state")
    if provider_state == "external" and not provider.get("name"):
        errors.append(f"{candidate_id}: external Provider State requires a name")
    if provider_state in {"internal", "unknown"} and provider.get("name") is not None:
        errors.append(f"{candidate_id}: {provider_state} Provider State cannot have a name")
    if provider_state == "unknown" and provider.get("roles"):
        errors.append(f"{candidate_id}: Unknown Provider State cannot have Provider Roles")

    field_evidence = candidate.get("field_evidence")
    if not isinstance(field_evidence, list):
        errors.append(f"{candidate_id}: field_evidence must be an array")
        return
    by_field: dict[str, list[str]] = {}
    for link in field_evidence:
        if not isinstance(link, dict):
            errors.append(f"{candidate_id}: field evidence link must be an object")
            continue
        field_path = link.get("field_path")
        if isinstance(field_path, str):
            if field_path in by_field:
                errors.append(f"{candidate_id}: duplicate evidence path {field_path!r}")
            by_field[field_path] = link.get("evidence_ids")
            _verify_evidence_references(
                f"{candidate_id}.{field_path}",
                link.get("evidence_ids"),
                evidence,
                errors,
            )

    required_fields = {"purchasing_status", "provider.state"}
    for index, concept in enumerate(concepts):
        required_fields.update(
            {
                f"concepts[{index}].normalized_name",
                f"concepts[{index}].observed_name_text",
                f"concepts[{index}].category_path",
            }
        )
        for key in ("quantity", "unit", "component_unit_price"):
            if concept.get(key) is not None:
                required_fields.add(f"concepts[{index}].{key}")
    if provider.get("name") is not None:
        required_fields.add("provider.name")
    if provider.get("observed_provider_text") is not None:
        required_fields.add("provider.observed_provider_text")
    if provider.get("roles"):
        required_fields.add("provider.roles")
    for key in (
        "commercial_quantity",
        "commercial_unit",
        "currency",
        "unit_price",
        "total_price",
        "purchase_date",
    ):
        if candidate.get(key) is not None:
            required_fields.add(key)
    missing = sorted(required_fields - set(by_field))
    if missing:
        errors.append(f"{candidate_id}: important fields lack evidence: {', '.join(missing)}")


def _verify_evidence_references(
    label: str,
    evidence_ids: Any,
    evidence: dict[str, dict[str, Any]],
    errors: list[str],
) -> None:
    if not isinstance(evidence_ids, list) or not evidence_ids:
        errors.append(f"{label}: evidence_ids must be non-empty")
        return
    unknown = sorted(set(evidence_ids) - set(evidence))
    if unknown:
        errors.append(f"{label}: unknown evidence ids {unknown}")


def _list(
    value: dict[str, Any],
    key: str,
    errors: list[str],
) -> list[Any]:
    result = value.get(key)
    if isinstance(result, list):
        return result
    errors.append(f"{key} must be an array")
    return []


def _used_range(sheet) -> str | None:
    non_empty = _non_empty_coordinates(sheet)
    if not non_empty:
        return None
    return sheet.calculate_dimension()


def _non_empty_coordinates(sheet) -> set[str]:
    return {
        cell.coordinate
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    }


def _coordinates_in_range(
    sheet,
    range_value: Any,
    errors: list[str],
    label: str,
) -> set[str]:
    if not isinstance(range_value, str):
        errors.append(f"{label}: range must be a string")
        return set()
    try:
        min_col, min_row, max_col, max_row = range_boundaries(range_value)
    except (TypeError, ValueError):
        errors.append(f"{label}: invalid range {range_value!r}")
        return set()
    if min_row < 1 or min_col < 1 or max_row < min_row or max_col < min_col:
        errors.append(f"{label}: invalid range {range_value!r}")
        return set()
    if max_row > sheet.max_row or max_col > sheet.max_column:
        errors.append(f"{label}: range {range_value!r} exceeds worksheet bounds")
        return set()
    return {
        f"{get_column_letter(column)}{row}"
        for row in range(min_row, max_row + 1)
        for column in range(min_col, max_col + 1)
    }


def canonical_cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return format(Decimal(str(value)).normalize(), "f")
    return str(value)


def normalized_amount(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip().replace(",", "").replace("PHP", "").replace("₱", "")
    try:
        return format(Decimal(text).normalize(), "f")
    except InvalidOperation:
        return text.casefold()
