from copy import copy
from datetime import date
from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from sqlalchemy import select

from backend.app.processing.worker import run_once


def _create_project(client):
    return client.post(
        "/api/project-workspaces",
        json={
            "project_name": "Arnaiz Residence Renovation",
            "project_type": "Residential renovation",
            "location": "Makati City",
            "completion_year": 2025,
            "contractor_assigned": "Internal",
        },
    ).json()


def _upload(client, project_id: int, content: bytes):
    return client.post(
        f"/api/project-workspaces/{project_id}/source-files",
        files={
            "files": (
                "purchase-log.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    ).json()


def _workbook_bytes(configure=None) -> bytes:
    workbook = Workbook()
    if configure is not None:
        configure(workbook)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class NoUsableRegionsProvider:
    provider_name = "fake"
    model = "fake-xlsx-model"

    def profile_worksheet(self, *, worksheet: dict, source_submission_id: int):
        return {
            "worksheet_name": worksheet["worksheet"]["name"],
            "title_rows": [],
            "header_rows": [],
            "regions": [],
        }


class ExplicitAnnotationColumnsProvider:
    provider_name = "fake"
    model = "fake-xlsx-model"

    def profile_worksheet(self, *, worksheet: dict, source_submission_id: int):
        return {
            "worksheet_name": worksheet["worksheet"]["name"],
            "title_rows": [],
            "header_rows": [1],
            "regions": [
                {
                    "region_id": "purchases",
                    "usable": True,
                    "unusable_reason": None,
                    "header_row_numbers": [1],
                    "body_start_row": 2,
                    "body_end_row": 3,
                    "columns": {"line_type": "A"},
                }
            ],
        }

    def extract_worksheet_chunk(self, **kwargs):
        return {"candidates": []}


def test_explicit_xlsx_annotation_columns_create_exact_cell_grounded_proposals(
    client, monkeypatch, tmp_path
):
    monkeypatch.setenv("HABI_STORAGE_ROOT", str(tmp_path))

    def configure(workbook):
        sheet = workbook.active
        sheet.title = "Purchases"
        sheet.append(
            [
                "Line kind",
                "Material name",
                "Material category path",
                "Provider State",
                "Provider name",
                "Delivery Terms",
                "Material Warranty",
                "Provider Notes",
            ]
        )
        sheet.append(
            [
                "Material",
                "PVC pipe",
                "Plumbing / Pipes",
                "External",
                "ABC Trading",
                "Delivery included",
                "Five-year warranty",
                "Accredited distributor",
            ]
        )
        sheet.append(["Not a purchase line"])

    project = _create_project(client)
    submission = _upload(client, project["id"], _workbook_bytes(configure))

    assert run_once(
        client.app.state.session_factory,
        ai_provider=ExplicitAnnotationColumnsProvider(),
    ) == 1
    job = client.get(
        f"/api/project-workspaces/{project['id']}/processing-jobs/"
        f"{submission['processing_job']['id']}"
    ).json()["processing_job"]
    assert job["status"] == "review_ready", job
    review = client.get(
        f"/api/project-workspaces/{project['id']}/review-batches/{job['review_batch_id']}"
    ).json()

    assert review["candidates"][0]["proposed_payload"]["annotation_proposals"] == [
        {
            "proposal_id": "xlsx:Purchases:F2",
            "text": "Delivery included",
            "annotation_type": "delivery_terms",
            "target": "purchase_line",
            "source_excerpt": "Delivery included",
            "source_locator": {
                "kind": "xlsx_cell",
                "worksheet": "Purchases",
                "row": 2,
                "column": "F",
                "coordinate": "F2",
            },
            "provenance": "source_field",
        },
        {
            "proposal_id": "xlsx:Purchases:G2",
            "text": "Five-year warranty",
            "annotation_type": "warranty_terms",
            "target": "material",
            "source_excerpt": "Five-year warranty",
            "source_locator": {
                "kind": "xlsx_cell",
                "worksheet": "Purchases",
                "row": 2,
                "column": "G",
                "coordinate": "G2",
            },
            "provenance": "source_field",
        },
        {
            "proposal_id": "xlsx:Purchases:H2",
            "text": "Accredited distributor",
            "annotation_type": "general_qualifier",
            "target": "provider",
            "source_excerpt": "Accredited distributor",
            "source_locator": {
                "kind": "xlsx_cell",
                "worksheet": "Purchases",
                "row": 2,
                "column": "H",
                "coordinate": "H2",
            },
            "provenance": "source_field",
        },
    ]


def test_deterministic_xlsx_annotations_are_capped_and_diagnosed(client, monkeypatch, tmp_path):
    monkeypatch.setenv("HABI_STORAGE_ROOT", str(tmp_path))
    annotation_headers = [
        "Delivery Terms",
        "Payment Terms",
        "Validity Terms",
        "Validity",
        "Warranty Terms",
        "Warranty",
        "Availability Terms",
        "Availability",
        "Condition or Exclusion",
        "Conditions",
        "Exclusions",
        "Remarks or Terms",
        "Remarks",
        "Terms",
        "Notes",
        "Material Delivery Terms",
        "Material Payment Terms",
        "Material Validity",
        "Material Warranty",
        "Material Availability",
        "Material Conditions",
    ]

    def configure(workbook):
        sheet = workbook.active
        sheet.title = "Purchases"
        sheet.append(
            [
                "Line kind",
                "Material name",
                "Material category path",
                "Provider State",
                "Provider name",
                *annotation_headers,
                "Provider Notes",
            ]
        )
        sheet.append(
            [
                "Material",
                "PVC pipe",
                "Plumbing / Pipes",
                "External",
                "ABC Trading",
                *[f"Qualifier {index}" for index in range(1, 22)],
                "Call site supervisor tomorrow",
            ]
        )
        sheet.append(["Not a purchase line"])

    project = _create_project(client)
    submission = _upload(client, project["id"], _workbook_bytes(configure))

    assert run_once(
        client.app.state.session_factory,
        ai_provider=ExplicitAnnotationColumnsProvider(),
    ) == 1
    job = client.get(
        f"/api/project-workspaces/{project['id']}/processing-jobs/"
        f"{submission['processing_job']['id']}"
    ).json()["processing_job"]
    assert job["status"] == "review_ready", job
    review = client.get(
        f"/api/project-workspaces/{project['id']}/review-batches/{job['review_batch_id']}"
    ).json()
    candidate = review["candidates"][0]["proposed_payload"]

    assert len(candidate["annotation_proposals"]) == 20
    assert candidate["annotation_omitted_count"] == 1
    assert candidate["annotation_detected_count"] == 21
    assert all(
        proposal["source_excerpt"] != "Call site supervisor tomorrow"
        for proposal in candidate["annotation_proposals"]
    )
    assert job["diagnostics"]["dropped_annotation_count"] == 2
    assert job["diagnostics"]["dropped_annotation_reasons"] == {
        "annotation_limit": 1,
        "workflow_noise": 1,
    }
    assert job["diagnostics"]["warning_summary"] == (
        "Annotation extraction limit reached — 20 of 21 source-grounded annotation "
        "proposals were retained. Review the source and add any omitted qualifiers that matter."
    )


class ValidXlsxProvider:
    provider_name = "fake"
    model = "fake-xlsx-model"

    def __init__(self, *, source_file_id: int):
        self.source_file_id = source_file_id
        self.profile_calls = []
        self.extraction_calls = []

    def profile_worksheet(self, *, worksheet: dict, source_submission_id: int):
        self.profile_calls.append((worksheet, source_submission_id))
        return {
            "worksheet_name": worksheet["worksheet"]["name"],
            "title_rows": [],
            "header_rows": [1],
            "regions": [
                {
                    "region_id": "purchases",
                    "usable": True,
                    "unusable_reason": None,
                    "header_row_numbers": [1],
                    "body_start_row": 2,
                    "body_end_row": 3,
                    "columns": {
                        "name": "A",
                        "quantity": "B",
                        "unit": "C",
                        "price": "D",
                    },
                }
            ],
        }

    def extract_worksheet_chunk(
        self,
        *,
        profile: dict,
        region: dict,
        rows: list[dict],
        context_rows: list[dict],
        source_submission_id: int,
    ):
        self.extraction_calls.append(
            {
                "profile": profile,
                "region": region,
                "rows": rows,
                "context_rows": context_rows,
                "source_submission_id": source_submission_id,
            }
        )
        return {
            "candidates": [
                {
                    "line_type": "material",
                    "name": "PVC pipe",
                    "quantity": "20",
                    "unit": "pcs",
                    "price": "1500",
                    "currency": "PHP",
                    "currency_state": "source_stated",
                    "provider_name": "ABC Trading",
                    "purchase_date": None,
                    "remarks_or_terms": None,
                    "confidence": 0.94,
                    "category_suggestion": {
                        "top_level_category": "Plumbing",
                        "subcategory": "Pipes",
                    },
                    "annotation_proposals": [
                        {
                            "text": "Delivery within seven days",
                            "annotation_type": "delivery_terms",
                            "target": "purchase_line",
                            "source_excerpt": "Delivery within seven days",
                        }
                    ],
                    "evidence": {
                        "source_submission_id": source_submission_id,
                        "source_file_id": self.source_file_id,
                        "worksheet": "Purchases",
                        "region_id": "purchases",
                        "primary_body_row": 2,
                        "locators": [{"row": 2, "role": "body"}],
                    },
                }
            ]
        }


class ChunkRecordingProvider:
    provider_name = "fake"
    model = "fake-xlsx-model"

    def __init__(self):
        self.extraction_calls = []

    def profile_worksheet(self, *, worksheet: dict, source_submission_id: int):
        row_numbers = [cell["row"] for cell in worksheet["cells"]]
        return {
            "worksheet_name": worksheet["worksheet"]["name"],
            "title_rows": [],
            "header_rows": [1],
            "regions": [
                {
                    "region_id": "purchases",
                    "usable": True,
                    "unusable_reason": None,
                    "header_row_numbers": [1],
                    "body_start_row": 2,
                    "body_end_row": max(row_numbers),
                    "columns": {"name": "A"},
                }
            ],
        }

    def extract_worksheet_chunk(self, **kwargs):
        self.extraction_calls.append(kwargs)
        return {"candidates": []}


class InvalidLocatorProvider(ValidXlsxProvider):
    def extract_worksheet_chunk(self, **kwargs):
        result = super().extract_worksheet_chunk(**kwargs)
        result["candidates"][0]["evidence"]["locators"] = [
            {"row": 999, "role": "body"}
        ]
        result["candidates"][0]["evidence"]["primary_body_row"] = 999
        return result


class PartialFailureProvider(ValidXlsxProvider):
    def __init__(self, *, source_file_id: int):
        super().__init__(source_file_id=source_file_id)
        self.attempt_count = 0

    def extract_worksheet_chunk(self, **kwargs):
        self.attempt_count += 1
        if self.attempt_count == 2:
            raise RuntimeError("second extraction call failed")
        return super().extract_worksheet_chunk(**kwargs)


class IncompleteFixtureLogProvider:
    provider_name = "fake"
    model = "fake-xlsx-model"

    def __init__(self, *, source_file_id: int):
        self.source_file_id = source_file_id

    def profile_worksheet(self, *, worksheet: dict, source_submission_id: int):
        return {
            "worksheet_name": "Sheet1",
            "title_rows": [1],
            "header_rows": [6],
            "regions": [
                {
                    "region_id": "purchase_lines_table_rows_7_plus",
                    "usable": True,
                    "unusable_reason": None,
                    "header_row_numbers": [6],
                    "body_start_row": 7,
                    "body_end_row": 15,
                    "columns": {
                        "line_type": None,
                        "name": "E",
                        "quantity": "G",
                        "unit": "H",
                        "price": "I",
                        "currency": "K",
                        "provider_name": "S",
                        "purchase_date": "L",
                        "remarks_or_terms": None,
                    },
                }
            ],
        }

    def extract_worksheet_chunk(
        self,
        *,
        profile: dict,
        region: dict,
        rows: list[dict],
        context_rows: list[dict],
        source_submission_id: int,
        memory_context: dict,
    ):
        def bundled_candidate(
            *,
            row: int,
            material_name: str,
            material_category: str,
            service_name: str,
            service_category: str,
            quantity: str,
            price: str,
            purchase_date: str,
        ) -> dict:
            return {
                "linked_concepts": [
                    {
                        "concept_type": "material",
                        "name": material_name,
                        "category_suggestion": {
                            "top_level_category": material_category,
                            "subcategory": material_category,
                        },
                    },
                    {
                        "concept_type": "service",
                        "name": service_name,
                        "category_suggestion": {
                            "top_level_category": service_category,
                            "subcategory": service_category,
                        },
                    },
                ],
                "quantity": quantity,
                "unit": "m²",
                "price": price,
                "currency": "PHP",
                "currency_state": "source_stated",
                "provider_state": "internal",
                "provider_name": "Habi Build Co.",
                "provider_category_suggestion": None,
                "purchase_date": purchase_date,
                "remarks_or_terms": None,
                "confidence": 0.6,
                "evidence": {
                    "source_submission_id": source_submission_id,
                    "source_file_id": self.source_file_id,
                    "worksheet": "Sheet1",
                    "region_id": region["region_id"],
                    "primary_body_row": row,
                    "locators": [{"row": row, "role": "body"}],
                },
            }

        return {
            "candidates": [
                bundled_candidate(
                    row=7,
                    material_name="Gypsum board ceiling",
                    material_category="Finishes / Ceilings",
                    service_name="Ceiling installation",
                    service_category="Services / Installation / Ceilings",
                    quantity="120",
                    price="108000",
                    purchase_date="2026-05-03",
                ),
                bundled_candidate(
                    row=14,
                    material_name="Acoustic insulation",
                    material_category="Finishes / Insulation (new gate)",
                    service_name="Insulation installation",
                    service_category="Services / Installation / Insulation (new gate)",
                    quantity="50",
                    price="80000",
                    purchase_date="2026-05-20",
                ),
            ]
        }


class MalformedColumnFixtureLogProvider(IncompleteFixtureLogProvider):
    def profile_worksheet(self, *, worksheet: dict, source_submission_id: int):
        profile = super().profile_worksheet(
            worksheet=worksheet, source_submission_id=source_submission_id
        )
        profile["regions"][0]["columns"]["material_name"] = "M/N"
        return profile


def _configure_import_fixture_workbook(workbook):
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["Import Fixtures"])
    sheet.append(
        [
            "Reviewer-final fields are the import truth. Apply the rows in fixture-ID order; "
            "FX-08 is intentionally blocked until its three category paths are resolved."
        ]
    )
    sheet.append([])
    sheet.append(
        [
            "All quantity, price, currency, date, and evidence fields are shared Purchase Line "
            "facts. A bundled line must never allocate combined price or quantity between its "
            "linked concepts."
        ]
    )
    sheet.append([])
    sheet.append(
        [
            "Fixture ID",
            "Project key",
            "Source submission",
            "Evidence record",
            "Purchase Line description",
            "Line kind",
            "Quantity",
            "Unit",
            "Unit price",
            "Combined price",
            "Currency",
            "Purchase date",
            "Material name",
            "Material category path",
            "Service name",
            "Service category path",
            "AI Provider State",
            "Reviewer Final State",
            "Provider name",
            "Provider category path",
            "Expected display roles",
        ]
    )
    rows = [
        ["FX-01", "PRJ-A", "SS-001", "EV-001", "Supply & install gypsum ceiling", "Bundled", 120, "m²", 900, 108000, "PHP", date(2026, 5, 3), "Gypsum board ceiling", "Finishes / Ceilings", "Ceiling installation", "Services / Installation / Ceilings", "Unknown", "External", "Cebu Ceiling Works", "Providers / Specialty Contractors", "Materials | Services | Supply & install"],
        ["FX-02", "PRJ-A", "SS-001", "EV-002", "Gypsum ceiling materials top-up", "Standard · Material", 20, "m²", 900, 18000, "PHP", date(2026, 5, 3), "Gypsum board ceiling", "Finishes / Ceilings / Acoustic (conflicting suggestion)", None, None, "External", "External", "cebu ceiling works", "Providers / General (conflicting suggestion)", "Materials"],
        ["FX-03", "PRJ-A", "SS-002", "EV-003", "Interior paint — two rooms", "Standard · Material", 50, "L", 300, 15000, "PHP", date(2026, 5, 9), "Interior paint", "Finishes / Painting", None, None, "External", "External", "PaintPro Cebu", "Providers / General", "Materials"],
        ["FX-04", "PRJ-A", "SS-002", "EV-004", "Interior paint — corridor", "Standard · Material", 25, "L", 300, 7500, "PHP", date(2026, 5, 9), "Interior paint", "Finishes / Painting", None, None, "External", "External", "PaintPro Cebu", "Providers / General", "Materials"],
        ["FX-05", "PRJ-A", "SS-003", "EV-005", "Electrical testing", "Standard · Service", 1, "job", 5000, 5000, "PHP", date(2026, 5, 12), None, None, "Electrical testing", "Services / Testing", "Unknown", "External", "VoltCheck", "Providers / Testing", "Services"],
        ["FX-06", "PRJ-A", "SS-004", "EV-006", "Self-performed wall patching materials", "Standard · Material", 10, "bags", 250, 2500, "PHP", date(2026, 5, 14), "Wall patching compound", "Finishes / Repair", None, None, "External", "Internal", "HABI   BUILD CO.", None, "Materials"],
        ["FX-07", "PRJ-A", "SS-005", "EV-007", "Minor repair labour", "Standard · Service", 1, "job", 5500, 5500, "PHP", date(2026, 5, 16), None, None, "Minor repair", "Services / Repair", "Unknown", "Unknown", None, None, None],
        ["FX-08", "PRJ-A", "SS-006", "EV-008", "Supply & install acoustic insulation", "Bundled", 50, "m²", 1600, 80000, "PHP", date(2026, 5, 20), "Acoustic insulation", "Finishes / Insulation (new gate)", "Insulation installation", "Services / Installation / Insulation (new gate)", "External", "External", "Insulate PH", "Providers / General (new gate)", "Materials | Services | Supply & install"],
        ["FX-09", "PRJ-B", "SS-B001", "EV-B001", "Supply gypsum board ceiling", "Standard · Material", 10, "m²", 900, 9000, "PHP", date(2026, 5, 22), "Gypsum board ceiling", "Finishes / Ceilings", None, None, "External", "External", "Cebu Ceiling Works", "Providers / General", "Materials"],
    ]
    for row in rows:
        sheet.append(row)


def _process_incomplete_import_fixture(client, monkeypatch, tmp_path):
    monkeypatch.setenv("HABI_STORAGE_ROOT", str(tmp_path))
    project = client.post(
        "/api/project-workspaces",
        json={
            "project_name": "Cebu Office Fit-out",
            "project_type": "Commercial fit-out",
            "location": "Cebu City",
            "completion_year": 2026,
            "contractor_assigned": "Habi Build Co.",
        },
    ).json()
    submission = _upload(
        client,
        project["id"],
        _workbook_bytes(_configure_import_fixture_workbook),
    )
    provider = IncompleteFixtureLogProvider(
        source_file_id=submission["source_file"]["id"]
    )

    assert run_once(client.app.state.session_factory, ai_provider=provider) == 1

    job = client.get(
        f"/api/project-workspaces/{project['id']}/processing-jobs/"
        f"{submission['processing_job']['id']}"
    ).json()["processing_job"]
    return project, submission, job


def test_malformed_ai_column_mapping_is_recovered_from_exact_xlsx_headers(
    client, monkeypatch, tmp_path
):
    monkeypatch.setenv("HABI_STORAGE_ROOT", str(tmp_path))
    project = client.post(
        "/api/project-workspaces",
        json={
            "project_name": "Cebu Office Fit-out",
            "project_type": "Commercial fit-out",
            "location": "Cebu City",
            "completion_year": 2026,
            "contractor_assigned": "Habi Build Co.",
        },
    ).json()
    submission = _upload(
        client,
        project["id"],
        _workbook_bytes(_configure_import_fixture_workbook),
    )
    provider = MalformedColumnFixtureLogProvider(
        source_file_id=submission["source_file"]["id"]
    )

    assert run_once(client.app.state.session_factory, ai_provider=provider) == 1

    job = client.get(
        f"/api/project-workspaces/{project['id']}/processing-jobs/"
        f"{submission['processing_job']['id']}"
    ).json()["processing_job"]
    assert job["status"] == "review_ready"
    assert job["candidate_count"] == 9
    assert job["diagnostics"]["invalid_profile_column_mapping_count"] == 1


def test_explicit_xlsx_purchase_rows_are_not_silently_omitted_from_review(
    client, monkeypatch, tmp_path
):
    _, _, job = _process_incomplete_import_fixture(client, monkeypatch, tmp_path)

    assert job["status"] == "review_ready"
    assert job["candidate_count"] == 9
    assert job["diagnostics"]["raw_candidate_count"] == 2
    assert job["diagnostics"]["valid_candidate_count"] == 9
    assert job["diagnostics"]["source_grounded_candidate_count"] == 9
    assert job["diagnostics"]["ai_candidate_replaced_count"] == 2


def test_explicit_xlsx_purchase_rows_override_ungrounded_ai_fields(
    client, monkeypatch, tmp_path
):
    project, _, job = _process_incomplete_import_fixture(client, monkeypatch, tmp_path)
    review = client.get(
        f"/api/project-workspaces/{project['id']}/review-batches/"
        f"{job['review_batch_id']}"
    ).json()
    candidates_by_row = {
        candidate["proposed_payload"]["evidence"]["primary_body_row"]: candidate[
            "proposed_payload"
        ]
        for candidate in review["candidates"]
    }

    assert {
        row: [concept["concept_type"] for concept in payload["linked_concepts"]]
        for row, payload in candidates_by_row.items()
    } == {
        7: ["material", "service"],
        8: ["material"],
        9: ["material"],
        10: ["material"],
        11: ["service"],
        12: ["material"],
        13: ["service"],
        14: ["material", "service"],
        15: ["material"],
    }
    assert (
        candidates_by_row[7]["provider_state"],
        candidates_by_row[7]["provider_name"],
    ) == ("external", "Cebu Ceiling Works")
    assert (
        candidates_by_row[14]["provider_state"],
        candidates_by_row[14]["provider_name"],
    ) == ("external", "Insulate PH")
    assert (
        candidates_by_row[12]["provider_state"],
        candidates_by_row[12]["provider_name"],
    ) == ("internal", "HABI   BUILD CO.")
    assert (
        candidates_by_row[13]["provider_state"],
        candidates_by_row[13]["provider_name"],
    ) == ("unknown", None)
    assert candidates_by_row[7]["price"] == "108000"
    assert candidates_by_row[7]["linked_concepts"][0]["category_suggestion"] == {
        "top_level_category": "Finishes",
        "subcategory": "Ceilings",
    }
    assert candidates_by_row[7]["linked_concepts"][1]["category_suggestion"] == {
        "top_level_category": "Services",
        "subcategory": "Installation / Ceilings",
    }


def test_corrupt_xlsx_is_preserved_and_processing_fails_without_review_batch(
    client, monkeypatch, tmp_path
):
    monkeypatch.setenv("HABI_STORAGE_ROOT", str(tmp_path))
    project = _create_project(client)
    submission = _upload(client, project["id"], b"not-an-openxml-workbook")

    assert run_once(client.app.state.session_factory, ai_provider=object()) == 1

    job = client.get(
        f"/api/project-workspaces/{project['id']}/processing-jobs/"
        f"{submission['processing_job']['id']}"
    ).json()["processing_job"]
    assert job["status"] == "failed"
    assert job["review_batch_id"] is None
    assert job["candidate_count"] == 0
    assert "workbook" in job["error_message"].lower()
    stored_path = tmp_path / submission["source_file"]["storage_path"]
    assert stored_path.read_bytes() == b"not-an-openxml-workbook"


def test_xlsx_processing_retains_visible_worksheet_artifact_and_skip_diagnostics(
    client, monkeypatch, tmp_path
):
    from backend.app.xlsx.models import WorksheetArtifact

    monkeypatch.setenv("HABI_STORAGE_ROOT", str(tmp_path))
    project = _create_project(client)

    def configure(workbook):
        visible = workbook.active
        visible.title = "Purchases"
        visible.append(["Item", "Price", "Hidden Column"])
        visible.append(["PVC pipe", 1500, "ignore me"])
        visible.append(["Hidden row", 999, None])
        visible.row_dimensions[3].hidden = True
        visible.column_dimensions["C"].hidden = True
        hidden = workbook.create_sheet("Internal Notes")
        hidden.sheet_state = "hidden"
        hidden["A1"] = "ignore this sheet"

    submission = _upload(client, project["id"], _workbook_bytes(configure))

    assert run_once(
        client.app.state.session_factory, ai_provider=NoUsableRegionsProvider()
    ) == 1

    job = client.get(
        f"/api/project-workspaces/{project['id']}/processing-jobs/"
        f"{submission['processing_job']['id']}"
    ).json()["processing_job"]
    assert job["status"] == "no_candidates_found"
    assert job["diagnostics"]["visible_worksheet_count"] == 1
    assert job["diagnostics"]["hidden_worksheet_count"] == 1
    assert job["diagnostics"]["hidden_row_count"] == 1
    assert job["diagnostics"]["hidden_column_count"] == 1
    assert job["diagnostics"]["artifact_count"] == 1

    with client.app.state.session_factory() as session:
        artifacts = session.scalars(select(WorksheetArtifact)).all()
        assert len(artifacts) == 1
        artifact = artifacts[0]
        assert artifact.worksheet_name == "Purchases"
        artifact_content = json.loads((tmp_path / artifact.artifact_path).read_text())

    assert [cell["coordinate"] for cell in artifact_content["cells"]] == [
        "A1",
        "B1",
        "A2",
        "B2",
    ]


def test_xlsx_processing_rejects_post_open_workbook_limits_without_artifacts(
    client, monkeypatch, tmp_path
):
    from backend.app.xlsx.models import WorksheetArtifact

    monkeypatch.setenv("HABI_STORAGE_ROOT", str(tmp_path))
    monkeypatch.setenv("HABI_XLSX_MAX_NON_EMPTY_ROWS", "1")
    project = _create_project(client)

    def configure(workbook):
        sheet = workbook.active
        sheet.append(["Item"])
        sheet.append(["PVC pipe"])

    submission = _upload(client, project["id"], _workbook_bytes(configure))
    run_once(client.app.state.session_factory, ai_provider=NoUsableRegionsProvider())
    job = client.get(
        f"/api/project-workspaces/{project['id']}/processing-jobs/"
        f"{submission['processing_job']['id']}"
    ).json()["processing_job"]

    assert job["status"] == "failed"
    assert "rows limit" in job["error_message"]
    assert job["diagnostics"]["visible_non_empty_row_count"] == 2
    assert job["review_batch_id"] is None
    with client.app.state.session_factory() as session:
        assert session.scalars(select(WorksheetArtifact)).all() == []


def test_worksheet_artifact_retains_formula_merge_format_and_style_metadata(
    client, monkeypatch, tmp_path
):
    from backend.app.xlsx.models import WorksheetArtifact

    monkeypatch.setenv("HABI_STORAGE_ROOT", str(tmp_path))
    project = _create_project(client)

    def configure(workbook):
        sheet = workbook.active
        sheet.title = "Purchases"
        sheet.merge_cells("A1:D1")
        sheet["A1"] = "Purchase Log"
        bold_font = copy(sheet["A1"].font)
        bold_font.bold = True
        sheet["A1"].font = bold_font
        sheet["A1"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
        sheet.append(["Item", "Qty", "Unit", "Price", "Total"])
        sheet.append(["PVC pipe", 2, "pcs", 1500, "=B3*D3"])
        sheet["D3"].number_format = "#,##0.00"

    submission = _upload(client, project["id"], _workbook_bytes(configure))
    run_once(client.app.state.session_factory, ai_provider=NoUsableRegionsProvider())

    with client.app.state.session_factory() as session:
        artifact = session.scalar(
            select(WorksheetArtifact).where(
                WorksheetArtifact.source_file_id == submission["source_file"]["id"]
            )
        )
        assert artifact is not None
        content = json.loads((tmp_path / artifact.artifact_path).read_text())

    assert content["worksheet"]["merged_ranges"] == ["A1:D1"]
    by_coordinate = {cell["coordinate"]: cell for cell in content["cells"]}
    assert by_coordinate["A1"]["is_bold"] is True
    assert by_coordinate["A1"]["has_fill"] is True
    assert by_coordinate["D3"]["raw_value"] == 1500
    assert by_coordinate["D3"]["normalized_value"] == 1500
    assert by_coordinate["D3"]["displayed_text"] == "1500"
    assert by_coordinate["D3"]["cell_type"] == "number"
    assert by_coordinate["D3"]["number_format"] == "#,##0.00"
    assert by_coordinate["E3"]["formula"] == "=B3*D3"
    assert by_coordinate["E3"]["raw_value"] is None
    assert by_coordinate["E3"]["cell_type"] == "formula"


def test_xlsx_processing_profiles_then_extracts_verified_candidates(
    client, monkeypatch, tmp_path
):
    monkeypatch.setenv("HABI_STORAGE_ROOT", str(tmp_path))
    project = _create_project(client)

    def configure(workbook):
        sheet = workbook.active
        sheet.title = "Purchases"
        sheet.append(["Item", "Qty", "Unit", "Price", "Additional detail"])
        sheet.append(["PVC pipe", 20, "pcs", 1500, "Delivery within seven days"])
        sheet.append(["Cement", 10, "bags", 2800])

    submission = _upload(client, project["id"], _workbook_bytes(configure))
    provider = ValidXlsxProvider(source_file_id=submission["source_file"]["id"])

    assert run_once(client.app.state.session_factory, ai_provider=provider) == 1

    job_detail = client.get(
        f"/api/project-workspaces/{project['id']}/processing-jobs/"
        f"{submission['processing_job']['id']}"
    ).json()
    assert job_detail["processing_job"]["status"] == "review_ready"
    assert job_detail["processing_job"]["candidate_count"] == 1
    assert len(provider.profile_calls) == 1
    assert len(provider.extraction_calls) == 1
    extraction_call = provider.extraction_calls[0]
    assert extraction_call["region"]["source_file_id"] == submission["source_file"]["id"]
    assert extraction_call["region"]["worksheet"] == "Purchases"
    assert [row["row"] for row in extraction_call["rows"]] == [2, 3]
    assert [row["row"] for row in extraction_call["context_rows"]] == [1]
    review = client.get(
        f"/api/project-workspaces/{project['id']}/review-batches/"
        f"{job_detail['review_batch_id']}"
    ).json()
    candidate = review["candidates"][0]
    assert candidate["source_file"]["original_filename"] == "purchase-log.xlsx"
    assert candidate["proposed_payload"]["name"] == "PVC pipe"
    assert candidate["proposed_payload"]["evidence"] == {
        "source_submission_id": submission["source_submission"]["id"],
        "source_file_id": submission["source_file"]["id"],
        "worksheet": "Purchases",
        "region_id": "purchases",
        "primary_body_row": 2,
        "locators": [{"row": 2, "role": "body"}],
    }
    assert candidate["proposed_payload"]["annotation_proposals"] == [
        {
            "proposal_id": "ai:xlsx:annotation:0",
            "text": "Delivery within seven days",
            "annotation_type": "delivery_terms",
            "target": "purchase_line",
            "source_excerpt": "Delivery within seven days",
            "source_locator": {
                "kind": "xlsx_cell",
                "worksheet": "Purchases",
                "row": 2,
                "column": 5,
                "coordinate": "E2",
            },
            "provenance": "ai_suggested",
        }
    ]


def test_xlsx_extraction_chunks_are_independent_and_respect_configured_row_limit(
    client, monkeypatch, tmp_path
):
    monkeypatch.setenv("HABI_STORAGE_ROOT", str(tmp_path))
    monkeypatch.setenv("HABI_XLSX_EXTRACTION_CHUNK_ROWS", "2")
    project = _create_project(client)

    def configure(workbook):
        sheet = workbook.active
        sheet.title = "Purchases"
        sheet.append(["Item"])
        for index in range(5):
            sheet.append([f"Item {index}"])

    submission = _upload(client, project["id"], _workbook_bytes(configure))
    provider = ChunkRecordingProvider()
    run_once(client.app.state.session_factory, ai_provider=provider)
    job = client.get(
        f"/api/project-workspaces/{project['id']}/processing-jobs/"
        f"{submission['processing_job']['id']}"
    ).json()["processing_job"]

    assert job["status"] == "no_candidates_found"
    assert [len(call["rows"]) for call in provider.extraction_calls] == [2, 2, 1]
    assert all(
        [row["row"] for row in call["context_rows"]] == [1]
        for call in provider.extraction_calls
    )
    assert job["diagnostics"]["extraction_chunk_count"] == 3


def test_xlsx_candidate_with_unverified_locator_is_dropped_before_review(
    client, monkeypatch, tmp_path
):
    monkeypatch.setenv("HABI_STORAGE_ROOT", str(tmp_path))
    project = _create_project(client)

    def configure(workbook):
        sheet = workbook.active
        sheet.title = "Purchases"
        sheet.append(["Item", "Qty", "Unit", "Price"])
        sheet.append(["PVC pipe", 20, "pcs", 1500])
        sheet.append(["Cement", 10, "bags", 2800])

    submission = _upload(client, project["id"], _workbook_bytes(configure))
    provider = InvalidLocatorProvider(source_file_id=submission["source_file"]["id"])
    run_once(client.app.state.session_factory, ai_provider=provider)
    job = client.get(
        f"/api/project-workspaces/{project['id']}/processing-jobs/"
        f"{submission['processing_job']['id']}"
    ).json()["processing_job"]

    assert job["status"] == "no_candidates_found"
    assert job["review_batch_id"] is None
    assert job["diagnostics"]["raw_candidate_count"] == 1
    assert job["diagnostics"]["valid_candidate_count"] == 0
    assert job["diagnostics"]["dropped_candidate_count"] == 1


def test_xlsx_technical_failure_creates_no_partial_review_output_and_retains_artifacts(
    client, monkeypatch, tmp_path
):
    from backend.app.review.models import ExtractedCandidate, ReviewBatch
    from backend.app.xlsx.models import WorksheetArtifact

    monkeypatch.setenv("HABI_STORAGE_ROOT", str(tmp_path))
    monkeypatch.setenv("HABI_XLSX_EXTRACTION_CHUNK_ROWS", "1")
    project = _create_project(client)

    def configure(workbook):
        sheet = workbook.active
        sheet.title = "Purchases"
        sheet.append(["Item", "Qty", "Unit", "Price"])
        sheet.append(["PVC pipe", 20, "pcs", 1500])
        sheet.append(["Cement", 10, "bags", 2800])

    submission = _upload(client, project["id"], _workbook_bytes(configure))
    provider = PartialFailureProvider(source_file_id=submission["source_file"]["id"])
    run_once(client.app.state.session_factory, ai_provider=provider)
    job = client.get(
        f"/api/project-workspaces/{project['id']}/processing-jobs/"
        f"{submission['processing_job']['id']}"
    ).json()["processing_job"]

    assert job["status"] == "failed"
    assert job["review_batch_id"] is None
    assert "second extraction call failed" in job["error_message"]
    with client.app.state.session_factory() as session:
        assert session.scalars(select(ReviewBatch)).all() == []
        assert session.scalars(select(ExtractedCandidate)).all() == []
        artifact = session.scalar(select(WorksheetArtifact))
        assert artifact is not None
        assert artifact.profile is not None
        assert (tmp_path / artifact.artifact_path).exists()


def test_unexpected_xlsx_processor_error_marks_job_failed(
    client, monkeypatch, tmp_path
):
    import backend.app.processing.worker as worker

    monkeypatch.setenv("HABI_STORAGE_ROOT", str(tmp_path))
    project = _create_project(client)
    submission = _upload(client, project["id"], b"preserved-source")

    def raise_unexpected(*args, **kwargs):
        raise RuntimeError("unexpected processor failure")

    monkeypatch.setattr(worker, "process_xlsx_source_file", raise_unexpected)

    assert worker.run_once(client.app.state.session_factory, ai_provider=object()) == 1
    job = client.get(
        f"/api/project-workspaces/{project['id']}/processing-jobs/"
        f"{submission['processing_job']['id']}"
    ).json()["processing_job"]
    assert job["status"] == "failed"
    assert job["finished_at"] is not None
    assert job["error_message"] == "unexpected processor failure"


def test_approved_xlsx_candidate_imports_verified_spreadsheet_evidence(
    client, monkeypatch, tmp_path
):
    from backend.app.evidence.models import EvidenceRecord

    monkeypatch.setenv("HABI_STORAGE_ROOT", str(tmp_path))
    project = _create_project(client)

    def configure(workbook):
        sheet = workbook.active
        sheet.title = "Purchases"
        sheet.append(["Item", "Qty", "Unit", "Price"])
        sheet.append(["PVC pipe", 20, "pcs", 1500])
        sheet.append(["Cement", 10, "bags", 2800])

    submission = _upload(client, project["id"], _workbook_bytes(configure))
    provider = ValidXlsxProvider(source_file_id=submission["source_file"]["id"])
    run_once(client.app.state.session_factory, ai_provider=provider)
    job_detail = client.get(
        f"/api/project-workspaces/{project['id']}/processing-jobs/"
        f"{submission['processing_job']['id']}"
    ).json()
    review = client.get(
        f"/api/project-workspaces/{project['id']}/review-batches/"
        f"{job_detail['review_batch_id']}"
    ).json()
    candidate = review["candidates"][0]

    draft = client.put(
        f"/api/project-workspaces/{project['id']}/review-batches/"
        f"{job_detail['review_batch_id']}/review-draft",
        json={
            "candidates": [
                {
                    "candidate_id": candidate["id"],
                    "included": True,
                    "reviewed_payload": {
                        "line_type": "material",
                        "name": "PVC pipe",
                        "top_level_category": "Plumbing",
                        "subcategory": "Pipes",
                        "quantity": "20",
                        "unit": "pcs",
                        "price": "1500",
                        "currency": "PHP",
                        "provider_name": "ABC Trading",
                    },
                }
            ]
        },
    )
    imported = client.post(
        f"/api/project-workspaces/{project['id']}/review-batches/"
        f"{job_detail['review_batch_id']}/import"
    )

    assert draft.status_code == 200
    assert imported.status_code == 200
    purchase_lines = client.get(
        f"/api/project-workspaces/{project['id']}/purchase-lines"
    ).json()["items"]
    assert purchase_lines[0]["source_label"] == "purchase-log.xlsx"
    with client.app.state.session_factory() as session:
        evidence = session.scalar(select(EvidenceRecord))
        assert evidence is not None
        assert evidence.manual_source_entry_id is None
        assert evidence.source_file_id == submission["source_file"]["id"]
        assert evidence.source_label == "purchase-log.xlsx"
        assert {
            key: evidence.content[key]
            for key in candidate["proposed_payload"]["evidence"]
        } == candidate["proposed_payload"]["evidence"]
        assert evidence.content["row_snapshot"] == [
            {
                "column": 1,
                "coordinate": "A2",
                "header": "Item",
                "value": "PVC pipe",
                "annotation": False,
            },
            {
                "column": 2,
                "coordinate": "B2",
                "header": "Qty",
                "value": "20",
                "annotation": False,
            },
            {
                "column": 3,
                "coordinate": "C2",
                "header": "Unit",
                "value": "pcs",
                "annotation": False,
            },
            {
                "column": 4,
                "coordinate": "D2",
                "header": "Price",
                "value": "1500",
                "annotation": False,
            },
        ]
